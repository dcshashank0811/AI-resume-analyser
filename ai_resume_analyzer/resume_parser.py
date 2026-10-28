from io import BytesIO

from pypdf import PdfReader
from docx import Document

MAX_MB = 5


def validate_file(uploaded_file) -> str | None:
    """Return an error message, or None if the file is fine."""
    name = uploaded_file.name.lower()
    if not name.endswith((".pdf", ".docx")):
        return "Only PDF and DOCX files are supported."
    if uploaded_file.size > MAX_MB * 1024 * 1024:
        return f"File is larger than {MAX_MB} MB."
    return None


def extract_text(uploaded_file) -> str:
    """Read text in memory only - nothing is written to disk."""
    data = BytesIO(uploaded_file.getvalue())
    if uploaded_file.name.lower().endswith(".pdf"):
        reader = PdfReader(data)
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    doc = Document(data)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)
